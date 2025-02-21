import re
import os
import glob
import time
import string
import warnings
import docx
import mysql.connector
from docx import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_openai import ChatOpenAI
from langchain.prompts import ChatPromptTemplate  
from langchain.chains import LLMChain   
from langchain_community.document_loaders import UnstructuredWordDocumentLoader
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation

# 示例：根据生成文本解析出题目、选项和答案
def parse_questions(text):
    """
    解析生成的文本，返回一个包含元组 (question_index, content, options, answer) 的列表。
    
    适用于格式：
    
    Q1 かれが　手伝って　（  　　　　　 ）　宿題 (しゅくだい) が　終わらなっかった。
    1　もらったから
    2　くれなかったから
    3　ほしいから
    4　ほしかったから
    Answer: 2

    Q2 うちの　子どもは　勉強 (べんきょう) しないで　（  　　　　　 ）　ばかりいる。
    1　あそび
    2　あそぶ
    3　あそばない
    4　あそんで
    Answer: 4
    """
    question_pattern = re.compile(r'(Q\d+.*?)\n(?:Answer:|$)\s*(\d)?', re.DOTALL)
    question_matches = question_pattern.findall(text)

    qa_list = []
    for q_text, answer in question_matches:
        lines = q_text.strip().splitlines()
        if not lines:
            continue

        # 解析题号和题目内容
        m = re.match(r'(Q\d+)[\.\s]*(.*)', lines[0].strip())
        if m:
            q_index = m.group(1)
            content = m.group(2).strip()
        else:
            q_index = ""
            content = lines[0].strip()

        # 解析选项
        options = []
        for line in lines[1:]:
            line = line.strip()
            m_opt = re.match(r'^(\d+)[\.\、\s]+(.*)', line)
            if m_opt:
                options.append(f"{m_opt.group(1)}. {m_opt.group(2).strip()}")
            else:
                # 处理题目可能换行的情况
                content += " " + line

        # 存入解析结果
        qa_list.append((q_index, content, "\n".join(options), answer))

    return qa_list

def store_questions_to_excel(qa_list, output, filename):
    """
    将题目信息存储到 Excel 文件中。
    """
    # 确保保存目录存在
    if not os.path.exists(output):
        os.makedirs(output)

    # 清理文件名中的非法字符
    filename = re.sub(r'[\\/*?:"<>|]', "", filename)

    # 创建 Excel 工作簿和工作表
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"

    # 设置表头
    headers = ["Question Index", "Content", "Options", "Answer", "Suggestions", "Modifications (if any)"]
    ws.append(headers)

    # 将题目信息写入表格
    for qa in qa_list:
        q_index, content, options, answer = qa
        if not options:
            options = "No options"
        if not answer:
            answer = "No answer"
        ws.append([q_index, content, options, answer, "", ""])

    # 设置各列宽度
    column_widths = {
        "A": 15,
        "B": 50,
        "C": 30,
        "D": 20,
        "E": 20,
        "F": 40
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 设置所有单元格自动换行
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True)

    # 创建下拉选择的数据验证
    dv = DataValidation(type="list", formula1='"High_Q,Low_Q,Drop,Minor changes"', allow_blank=True)
    ws.add_data_validation(dv)
    dv_range = f"E2:E{ws.max_row}"
    dv.add(dv_range)

    # 保存 Excel 文件
    output_path = os.path.join(output, filename)
    print(f"Saving Excel file to: {output_path}")
    wb.save(output_path)
    print(f"Successfully stored in {output_path}")



def process_word_to_excel(doc_filepath, excel_output):
    """
    Process the content from each Word document in a folder and store them in separate Excel files.
    """
    # 检查文件夹是否存在
    if not os.path.exists(doc_filepath):
        print(f"Error: Directory not found - {doc_filepath}")
        return

    # 获取文件夹中所有的 .docx 文件
    doc_files = glob.glob(os.path.join(doc_filepath, "*.docx"))
    if not doc_files:
        print("No .docx files found in the directory.")
        return

    for filepath in doc_files:
        try:
            filename = os.path.splitext(os.path.basename(filepath))[0]
            doc = Document(filepath)
            all_text = "\n".join([para.text for para in doc.paragraphs])

            # 解析提取的文本
            qa_list = parse_questions(all_text)

            # 将解析的题目和答案保存到 Excel 文件，文件名与 Word 文件名一致
            excel_filename = f"{filename}.xlsx"
            store_questions_to_excel(qa_list, excel_output, excel_filename)

        except Exception as e:
            print(f"Error processing {filepath}: {e}")